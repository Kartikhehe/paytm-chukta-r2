"""Job D: assemble submission/ and print a pass/fail checklist.

  .venv/bin/python submission_check.py

Copies ONLY KartikRaj_R2.pdf and KartikRaj_R2_Annexure.xlsx from chukta_handoff/Chukta_R2_kit/out/
into a fresh submission/ folder, then checks: file names, 19 pages, 0 amber boxes, 0 formula errors
(LibreOffice recalc of a temporary copy; the submitted file is not modified), the prototype link
returns HTTP 200, banned wording absent from the PDF text (pdftotext), PDF under 10 MB.
"""
import os, re, shutil, subprocess, sys, tempfile, urllib.request
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(ROOT, 'chukta_handoff', 'Chukta_R2_kit', 'out')
SUB = os.path.join(ROOT, 'submission')
FILES = ['KartikRaj_R2.pdf', 'KartikRaj_R2_Annexure.xlsx']
URL = 'https://chukta-r2-prototype.vercel.app'
BANNED = [r'\bvideo\b', r'walkthrough\s+completed', r'paid\s+a\s+real\s+bill']
SOFFICE = shutil.which('soffice') or '/Applications/LibreOffice.app/Contents/MacOS/soffice'
ERR = re.compile(r'^#(REF!|VALUE!|DIV/0!|NAME\?|N/A|NUM!|NULL!)')

if os.path.isdir(SUB):
    shutil.rmtree(SUB)
os.makedirs(SUB)
for f in FILES:
    shutil.copy2(os.path.join(OUT, f), os.path.join(SUB, f))
pdf, xlsx = (os.path.join(SUB, f) for f in FILES)
rows = []

def check(name, ok, detail):
    rows.append((name, ok, detail))

check('Exact file names', sorted(os.listdir(SUB)) == sorted(FILES), ', '.join(sorted(os.listdir(SUB))))
n = len(PdfReader(pdf).pages)
check('19 pages (cover + 15 + 3 annexures)', n == 19, f'{n} pages')
html = open(os.path.join(OUT, 'deck.html'), encoding='utf-8').read()
tbd = html.count('class="tbd"')
check('0 amber boxes in deck.html', tbd == 0, f'{tbd} found')

tmp = tempfile.mkdtemp()
subprocess.run([SOFFICE, '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', tmp, xlsx],
               check=True, capture_output=True, timeout=240)
wv = load_workbook(os.path.join(tmp, FILES[1]), data_only=True)
wf = load_workbook(xlsx)
nform, errs = 0, []
for ws in wf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                nform += 1
                v = wv[ws.title][c.coordinate].value
                if isinstance(v, str) and ERR.match(v):
                    errs.append(f'{ws.title}!{c.coordinate}')
check('0 formula errors (LibreOffice recalc)', not errs, f'{nform} formulas, {len(errs)} errors {errs[:5]}')

try:
    code = urllib.request.urlopen(urllib.request.Request(URL, headers={'User-Agent': 'submission-check'}), timeout=20).status
except Exception as e:  # noqa: BLE001
    code = str(e)
check('Prototype link opens (HTTP 200)', code == 200, f'{URL} → {code}')

txt = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True, check=True).stdout
hits = {p: len(re.findall(p, txt, re.I)) for p in BANNED}
check('No "video" / "walkthrough completed" / "paid a real bill" in PDF text', not any(hits.values()), str(hits))
size = os.path.getsize(pdf) / 1e6
check('PDF under 10 MB', size < 10, f'{size:.2f} MB')

w = max(len(r[0]) for r in rows)
print()
for name, ok, detail in rows:
    print(f"{'PASS' if ok else 'FAIL'}  {name:<{w}}  {detail}")
print(f'\nsubmission/ → {SUB}')
sys.exit(0 if all(r[1] for r in rows) else 1)
