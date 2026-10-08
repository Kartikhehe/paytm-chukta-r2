"""Build KartikRaj_R2_Annexure.xlsx (formulas live; recalc with LibreOffice afterwards)."""
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
import model

A = model.A
NAVY, CYAN = '002E6E', '00BAF2'
F = 'Arial'
fT = Font(name=F, size=14, bold=True, color=NAVY)
fH = Font(name=F, size=10, bold=True, color='FFFFFF')
fB = Font(name=F, size=10)
fBold = Font(name=F, size=10, bold=True)
fIn = Font(name=F, size=10, color='0000FF')
fLink = Font(name=F, size=10, color='008000')
fNote = Font(name=F, size=9, italic=True, color='5B6B82')
fillH = PatternFill('solid', fgColor=NAVY)
fillSub = PatternFill('solid', fgColor='E3F5FD')
fillY = PatternFill('solid', fgColor='FFF9C4')
fillKey = PatternFill('solid', fgColor='FFFF00')
fillG = PatternFill('solid', fgColor='F3F7FB')
thin = Side(style='thin', color='D3DFEB')
bd = Border(top=thin, bottom=thin, left=thin, right=thin)
wrap = Alignment(wrap_text=True, vertical='top')

wb = Workbook()


def sheet(name, title, sub=None, widths=None):
    ws = wb.create_sheet(name)
    ws['A1'] = title; ws['A1'].font = fT
    if sub:
        ws['A2'] = sub; ws['A2'].font = fNote
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[L(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def header(ws, row, labels, col=1):
    for j, t in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=t)
        c.font = fH; c.fill = fillH; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = bd
    ws.row_dimensions[row].height = 30


def put(ws, ref, val, font=fB, fill=None, fmt=None, al=None):
    c = ws[ref]; c.value = val; c.font = font; c.border = bd
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    c.alignment = al or wrap
    return c


def rows(ws, start, data, fonts=None):
    for i, r in enumerate(data):
        for j, v in enumerate(r):
            c = ws.cell(row=start + i, column=1 + j, value=v)
            c.font = (fonts[j] if fonts else fB); c.border = bd; c.alignment = wrap


# ======================= READ ME =======================
ws = wb.active; ws.title = 'Read Me'; ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 34; ws.column_dimensions['C'].width = 100
ws['B1'] = 'Paytm Innovation Challenge 2026 · Round 2 Annexure · Paytm Chukta'; ws['B1'].font = fT
ws['B2'] = 'Kartik Raj · IIT Kanpur · Individual entry · Track B: Grow Paytm usage among merchants as UPI payers'; ws['B2'].font = fNote
info = [
    ('What this is', 'Supporting evidence for KartikRaj_R2.pdf. Round 1 evidence (50 VOCs) is copied in "R1 VOC" and remains in KartikRaj_VOC.xlsx.'),
    ('How Round 2 was run', 'In-person interviews in Kanpur on 8 Oct 2026: 21 shop owners (screened: decides supplier payments, receives on Paytm, paid a supplier in the last 7 days) and 4 distributors. Concept shown on a board; all answers given verbally and recorded as voice notes, then logged here.'),
    ('What the evidence is', 'Stated answers and stated commitments. No prototype walkthroughs or live payments were run in Round 2: usability and real payment behaviour are tested in week 1–2 of the 90-day pilot.'),
    ('Verbatims', 'Column N is transcribed from voice notes. Where several owners made the same point, it is logged under one common wording.'),
    ('R2 Validation Log', 'One row per shop owner. Ladder columns Q–T record what the owner SAID they would do, not observed behaviour.'),
    ('Distributor Log', 'One row per distributor / wholesaler (software used, payment mix, ledger offer, pilot interest).'),
    ('R2 Validation Dashboard', 'Every number is a formula on the two logs; these are the figures on slides 1, 6 and 7.'),
    ('Business Model', 'Opportunity sizing, Year 1–3 GMV, segment share, revenue, cost, cost per active shop and per ₹1,000 GMV. Blue = input (each with source); black = formula.'),
    ('Sensitivity', 'One-at-a-time tornado on the Year 3 drivers (slide 15).'),
    ('Pilot Design', '90-day cluster-randomised Kanpur pilot: cohort, targets, decision rules, power calculation (slide 12).'),
    ('Journey · Build · Risks', 'Before/after journey, build effort and dependencies, risk register (slides 8, 10, 15).'),
    ('Sources', 'Every external number with its link.'),
    ('Colour legend', 'Yellow fill = field data · Blue text = hard-coded input / assumption · Black = formula · Green = link to another sheet.'),
]
for i, (k, v) in enumerate(info, 4):
    put(ws, f'B{i}', k, fBold, fillSub); put(ws, f'C{i}', v)

# ======================= VALIDATION LOG =======================
cols = [('Respondent ID', 11), ('Date', 11), ('Area', 16), ('Category', 20), ('Role', 15), ('Receives customer payments on', 18),
        ('R1 respondent re-contacted?', 11), ('Last supplier payment ₹', 12), ('App used for it', 14), ('Paid from which phone', 15),
        ('Minutes taken (last real payment)', 11), ('A1 Usefulness 1–5', 10), ('A2 Move next bill to Chukta?', 11),
        ('A3 Main objection (from voice note)', 38), ('Objection code', 22), ('Bills they receive (stated)', 16),
        ('L1 Would share a real bill to be read (said)', 12), ('L2 Would add suppliers + 9 pm reminder (said)', 12),
        ('L3 Would pay next supplier bill via Paytm (said)', 13), ('L4 Wants to join 90-day pilot (said)', 11),
        ('Wants Chukta on personal phone', 11), ('Notes', 30)]
vl = sheet('R2 Validation Log', 'R2 Validation Log · one row per shop owner (screener + concept + stated ladder)',
           'Interviews 8 Oct 2026, Kanpur. Answers given verbally (voice notes). Anonymised: area + shop type only.', [w for _, w in cols])
header(vl, 5, [c for c, _ in cols])
N0, N1 = 6, 105
for r in range(N0, N1 + 1):
    for j in range(1, len(cols) + 1):
        c = vl.cell(row=r, column=j); c.fill = fillY; c.border = bd; c.font = fB
    vl.cell(row=r, column=2).number_format = 'dd-mmm-yy'
    vl.cell(row=r, column=8).number_format = '#,##0'
vl.freeze_panes = 'B6'
lists = {
    'D': 'Kirana / grocery,General / FMCG store,Pharmacy / chemist,Dairy / bakery / sweets,Restaurant / dhaba / canteen,Hardware / paint / sanitary,Mobile / electronics,Stationery / books,Clothing / footwear,Other',
    'E': 'Owner,Family member / partner,Manager who pays suppliers',
    'F': 'Paytm Soundbox,Paytm QR only,Paytm + other QR',
    'G': 'Yes,No', 'I': 'PhonePe,Google Pay,Paytm,BHIM,Bank app,Cash,NEFT / RTGS,Other',
    'J': 'Owner personal phone,Shop phone,Staff phone,Not by phone',
    'L': '1,2,3,4,5', 'M': 'Yes,Maybe,No',
    'O': 'Distributor must agree,Trust / fraud worry,Prefer current app,Too many steps,Phone / account mismatch,Bills are paper only,Wants credit first,Data privacy,Other,None',
    'P': 'e-Invoice QR,Paper / photo / WhatsApp,No bills',
    'Q': 'Yes,No', 'R': 'Yes,No', 'S': 'Yes,No,Already paid,No bill due', 'T': 'Yes,No', 'U': 'Yes,No'}
for col, opts in lists.items():
    dv = DataValidation(type='list', formula1=f'"{opts}"', allow_blank=True)
    vl.add_data_validation(dv); dv.add(f'{col}{N0}:{col}{N1}')
LOG = "'R2 Validation Log'!"
def rng(c): return f"{LOG}${c}${N0}:${c}${N1}"

# ======================= DISTRIBUTOR LOG =======================
dcols = [('Distributor ID', 11), ('Date', 11), ('Market / area', 18), ('Category', 16), ('Retailers served', 11), ('Monthly sales ₹ (approx)', 14),
         ('Retailer credit days', 10), ('Billing software', 13), ('Issues e-invoices (IRN)?', 11), ('Share of collections: cash %', 11),
         ('Share: UPI QR %', 10), ('Share: bank transfer %', 10), ('Concern about 0.4% MDR?', 12), ('Offered to share 4-week payment ledger?', 12),
         ('Agreed to join pilot (verbal, non-binding)?', 12), ('Key quote (from voice note)', 42), ('Notes', 26)]
dl = sheet('Distributor Log', 'Distributor Log · supply-side interviews (the actor whose salesman shows the QR)',
           'Interviews 8 Oct 2026, Kanpur. Owner or accounts head; answers given verbally.', [w for _, w in dcols])
header(dl, 5, [c for c, _ in dcols])
D0, D1 = 6, 25
for r in range(D0, D1 + 1):
    for j in range(1, len(dcols) + 1):
        c = dl.cell(row=r, column=j); c.fill = fillY; c.border = bd; c.font = fB
    dl.cell(row=r, column=2).number_format = 'dd-mmm-yy'
for col, opts in {'D': 'FMCG,Pharma,Dairy,Hardware,Electronics,Mixed', 'H': 'Tally,Marg,Busy,Other ERP,None / paper',
                  'I': 'Yes,No,Not sure', 'M': 'Yes (concerned),No / not aware', 'N': 'Yes,No,Maybe later', 'O': 'Yes,No,Maybe later'}.items():
    dv = DataValidation(type='list', formula1=f'"{opts}"', allow_blank=True); dl.add_data_validation(dv); dv.add(f'{col}{D0}:{col}{D1}')
DLG = "'Distributor Log'!"
def drng(c): return f"{DLG}${c}${D0}:${c}${D1}"

# ======================= DASHBOARD =======================
db = sheet('R2 Validation Dashboard', 'R2 Validation Dashboard · all formulas on the two logs',
           'Stated answers, 8 Oct 2026. The "deck key" column matches validation.json used to build the PDF.', [56, 14, 14, 22, 50])
header(db, 4, ['Metric', 'Value', 'Base (n)', 'deck key', 'How it is computed'])
NA = '"–"'
metrics = [
    ('Shop owners interviewed', f'=COUNTA({rng("A")})', '', 'merchants_tested', 'Rows with a Respondent ID'),
    ('Distributors interviewed', f'=COUNTA({drng("A")})', '', 'distributors_tested', 'Distributor Log rows'),
    ('Concept test respondents (n)', f'=COUNT({rng("L")})', '', 'concept_n', 'A1 answered'),
    ('Rated usefulness 4 or 5', f'=COUNTIF({rng("L")},">=4")', '=B7', 'concept_top2', 'A1'),
    ('Would move next bill to Chukta: Yes', f'=COUNTIF({rng("M")},"Yes")', '=B7', 'a2_yes', 'A2 forced choice vs current app'),
    ('Would move next bill to Chukta: Maybe', f'=COUNTIF({rng("M")},"Maybe")', '=B7', 'a2_maybe', 'A2'),
    ('Ladder respondents (n)', f'=COUNTA({rng("Q")})', '', 'walk_n', 'Answered the ladder questions'),
    ('Median minutes for their last real supplier payment', f'=IF(COUNT({rng("K")})=0,{NA},MEDIAN({rng("K")}))', '', 'today_min', 'Screener S3'),
    ('L1 Said they would share a real bill', f'=COUNTIF({rng("Q")},"Yes")', '=B11', 'l1_bill', 'Stated'),
    ('L2 Said they would add suppliers + reminder', f'=COUNTIF({rng("R")},"Yes")', '=B11', 'l2_reminder', 'Stated'),
    ('L3 Said they would pay next supplier bill via Paytm', f'=COUNTIF({rng("S")},"Yes")', '=B11', 'l3_paid', 'Stated, not observed'),
    ('L4 Want to join the 90-day pilot', f'=COUNTIF({rng("T")},"Yes")', '=B11', 'l4_pilot', 'Stated'),
    ('Want Chukta on personal phone', f'=COUNTIF({rng("U")},"Yes")', '=B11', 'personal_phone', ''),
    ('Shops that receive supplier bills (stated)', f'=COUNTA({rng("P")})-COUNTIF({rng("P")},"No bills")', '', 'bills_seen', ''),
    ('…whose bills carry an e-invoice QR (stated)', f'=COUNTIF({rng("P")},"e-Invoice QR")', '=B18', 'einv_bills', ''),
    ('Distributors who agreed to join the pilot', f'=COUNTIF({drng("O")},"Yes")', '=B6', 'd_loi', 'Verbal, non-binding'),
    ('Distributors who offered a 4-week ledger', f'=COUNTIF({drng("N")},"Yes")', '=B6', 'd_ledger', ''),
    ('Distributors on Tally / Marg / Busy', f'=COUNTIF({drng("H")},"Tally")+COUNTIF({drng("H")},"Marg")+COUNTIF({drng("H")},"Busy")', '=B6', 'd_erp', ''),
]
for i, (m, f, base, key, how) in enumerate(metrics, 5):
    put(db, f'A{i}', m); put(db, f'B{i}', f, fBold, al=Alignment(horizontal='right'))
    put(db, f'C{i}', base or None, fB, al=Alignment(horizontal='right')); put(db, f'D{i}', key, fNote); put(db, f'E{i}', how, fNote)
db.column_dimensions['F'].width = 12
put(db, 'F4', '% of base', fH, fillH)
for i in range(5, 5 + len(metrics)):
    put(db, f'F{i}', f'=IF(OR(C{i}="",C{i}=0),"",B{i}/C{i})', fB, fmt='0%', al=Alignment(horizontal='right'))
qr = 5 + len(metrics) + 1
put(db, f'A{qr}', 'Quota check (slide 6)', fH, fillH); put(db, f'B{qr}', 'Interviewed', fH, fillH); put(db, f'C{qr}', 'Quota', fH, fillH); put(db, f'D{qr}', 'deck key', fH, fillH)
quota = [('Kirana / general FMCG', f'=COUNTIF({rng("D")},"Kirana / grocery")+COUNTIF({rng("D")},"General / FMCG store")', 10, 'q_kirana'),
         ('Chemist / pharmacy', f'=COUNTIF({rng("D")},"Pharmacy / chemist")', 5, 'q_pharma'),
         ('Canteen / dairy / bakery', f'=COUNTIF({rng("D")},"Restaurant / dhaba / canteen")+COUNTIF({rng("D")},"Dairy / bakery / sweets")', 5, 'q_canteen'),
         ('Hardware / electronics / other', f'=B5-B{qr+1}-B{qr+2}-B{qr+3}', 5, 'q_other')]
for i, (a, f, q, k) in enumerate(quota, qr + 1):
    put(db, f'A{i}', a); put(db, f'B{i}', f, fBold); put(db, f'C{i}', q, fIn); put(db, f'D{i}', k, fNote)
db.column_dimensions['G'].width = 18; db.column_dimensions['H'].width = 12
put(db, f'G{qr}', 'L3 said yes (slide 7)', fH, fillH); put(db, f'H{qr}', 'deck key', fH, fillH)
def l3c(*cats): return '=' + '+'.join(f'COUNTIFS({rng("D")},"{c}",{rng("S")},"Yes")' for c in cats)
l3rows = [(l3c('Kirana / grocery', 'General / FMCG store'), 'l3_kirana'), (l3c('Pharmacy / chemist'), 'l3_pharma'),
          (l3c('Restaurant / dhaba / canteen', 'Dairy / bakery / sweets'), 'l3_canteen'), (f'=B15-G{qr+1}-G{qr+2}-G{qr+3}', 'l3_other')]
for i, (f, k) in enumerate(l3rows, qr + 1):
    put(db, f'G{i}', f, fBold); put(db, f'H{i}', k, fNote)
pb = qr + 6
put(db, f'A{pb}', 'Pass bars (slide 7)', fH, fillH); put(db, f'B{pb}', 'Result', fH, fillH); put(db, f'C{pb}', 'Bar', fH, fillH); put(db, f'D{pb}', 'Verdict', fH, fillH)
bars = [('H1 Desirability (stated): share saying they would pay next bill via Paytm', '=F15', 0.30),
        ('H2 Channel: distributors agreeing to join the pilot (count)', '=B20', 2)]
for i, (a, ref, bar) in enumerate(bars, pb + 1):
    put(db, f'A{i}', a)
    put(db, f'B{i}', ref, fBold, fmt='0%' if bar < 1 else '0')
    put(db, f'C{i}', bar, fIn, fmt='0%' if bar < 1 else '0')
    put(db, f'D{i}', f'=IF(B{i}="","–",IF(B{i}>=C{i},"PASS","BELOW BAR"))', fBold)
for i, a in enumerate(['H3 Feasibility: bills captured without typing', 'Usability: tasks completed unaided on the prototype'], pb + 3):
    put(db, f'A{i}', a); put(db, f'B{i}', 'Not run in R2', fB); put(db, f'C{i}', '70%', fIn); put(db, f'D{i}', 'Pilot week 1–2', fB)
ob = pb + 6
put(db, f'A{ob}', 'Objection codes (for "We heard → our response")', fH, fillH); put(db, f'B{ob}', 'Count', fH, fillH)
for i, code in enumerate(lists['O'].split(','), ob + 1):
    put(db, f'A{i}', code); put(db, f'B{i}', f'=COUNTIF({rng("O")},"{code}")', fBold)

# ======================= BUSINESS MODEL =======================
bm = sheet('Business Model', 'Business Model · Paytm Chukta (national roll-out)',
           'Blue = input with source. Black = formula. ₹ Cr unless stated. Same inputs generate the deck (slides 13–15, A1).', [52, 16, 16, 16, 16, 60])
header(bm, 4, ['A. Inputs (single value)', 'Value', '', '', '', 'Source / reasoning'])
inp = [('Paytm device merchants', A['devices'], '#,##0', 'Paytm Q1 FY27 results, Jul 2026 (1.57 crore device merchants)'),
       ('Share buying stock from suppliers', A['stock_share'], '0%', 'Assumption: retail-heavy Soundbox base; excludes services'),
       ('…of which pay suppliers over UPI monthly', A['upi_share'], '0%', 'Conservative; Kanpur R1 VOC: 48 of 50'),
       ('Supplier UPI per shop per month (₹)', A['spend'], '#,##0', '₹3 L sales × ~80% stock × ~25% UPI; Kanpur R1 average ₹1.09 L'),
       ('Paytm share of this UPI today', A['base_share'], '0.0%', 'Paytm national app share 8.1% (Aug 2026); Kanpur R1 sample 0%'),
       ('Supplier payments per shop per month (all apps)', A['bills'], '0', 'R1 VOC use-case frequencies'),
       ('MDR-eligible share of Chukta GMV (P2M > ₹2,000)', A['mdr_elig'], '0%', 'Rest: bank-a/c rail, P2PM-exempt small suppliers, tickets ≤ ₹2,000'),
       ('Payer-app (TPAP) share of MDR', A['tpap'], '0.00%', 'NPCI 15 Sep 2026: 0.4% MDR, 20% to UPI app'),
       ('Share of supplier QRs acquired by Paytm', A['acq_share'], '0%', 'Assumption: Chukta Collect onboards distributors to Paytm QR'),
       ('Acquirer share of MDR', A['acq'], '0.00%', 'NPCI 15 Sep 2026: 30% of 0.4% to acquirer'),
       ('Credit drawn per credit user per month (₹)', A['credit_amt'], '#,##0', 'Assumption: ~1/3 of monthly supplier UPI'),
       ('Paytm distribution take on credit', A['credit_take'], '0.0%', 'Assumption; depends on lender contracts'),
       ('Build cost, one-time (₹ Cr, Year 1)', A['build'], '0.0', 'Assumption: 1 squad × 12 weeks + partner integrations'),
       ('Run cost: team + infra (₹ Cr / yr)', A['run'], '0.0', 'Assumption'),
       ('CAC per new active shop (₹)', A['cac'], '#,##0', 'Build-up: 71 onboarding + 71 salesman fee + 43 training + 5 comms + 30 ops/fraud'),
       ('AI read cost per photo / PDF bill (₹)', A['capture'], '0.00', 'Assumption: vision model / OCR per page'),
       ('Share of bills needing AI read', A['ocr_share'], '0%', 'Rest arrive via e-invoice QR, ERP push or CSV'),
       ('Support cost per active shop per year (₹)', A['support'], '#,##0', 'Assumption'),
       ('National UPI value per month (₹ Cr)', A['upi_national'], '#,##0', 'NPCI, Aug 2026: ₹29.9 lakh crore')]
IN = {}
for i, (k, v, fm, src) in enumerate(inp, 5):
    put(bm, f'A{i}', k); put(bm, f'B{i}', v, fIn, fmt=fm, al=Alignment(horizontal='right')); put(bm, f'F{i}', src, fNote)
    IN[k] = f'$B${i}'
dev, stock, upi, spend, base, bills, elig, tpap, acqs, acq, camt, take, build, run, cac, cap, ocr, sup, nat = [f'$B${i}' for i in range(5, 24)]
r = 25
header(bm, r, ['B. Inputs by year', 'Year 1', 'Year 2', 'Year 3', '', 'Source / reasoning'])
for i, (k, vals, src) in enumerate([('Adoption: % of target shops active', A['adopt'], 'Cut from R1 (15/30/45%) after stress test; pilot activation bar 40% of reachable'),
                                    ('Paytm share of an adopter\'s supplier UPI', A['share'], 'Remainder: cash-only and non-Chukta suppliers'),
                                    ('Credit users among adopters', A['credit_users'], 'Credit launches after payment habit; lender-approved only')], r + 1):
    put(bm, f'A{i}', k)
    for j, v in enumerate(vals):
        c = put(bm, f'{"BCD"[j]}{i}', v, fIn, fmt='0%', al=Alignment(horizontal='right')); c.fill = fillKey
    put(bm, f'F{i}', src, fNote)
ADR, SHR, CUR = r + 1, r + 2, r + 3
r = 30
header(bm, r, ['C. Segment', 'Value', '', '', '', 'Formula'])
put(bm, 'A31', 'Target shops (device × stock × UPI)'); put(bm, 'B31', f'={dev}*{stock}*{upi}', fBold, fmt='#,##0'); put(bm, 'F31', 'Slide 13 funnel', fNote)
put(bm, 'A32', 'Pool: supplier UPI per month (₹ Cr)'); put(bm, 'B32', f'=B31*{spend}/10000000', fBold, fmt='#,##0')
put(bm, 'A33', 'Pool per year (₹ Cr)'); put(bm, 'B33', '=B32*12', fBold, fmt='#,##0')
put(bm, 'A34', 'Paytm GMV from pool today (₹ Cr / month)'); put(bm, 'B34', f'=B32*{base}', fBold, fmt='#,##0')
r = 36
header(bm, r, ['D. Year-by-year', 'Year 1', 'Year 2', 'Year 3', '3-yr total', 'Formula note'])
lines = [
    ('Active Chukta shops', lambda c: f'=$B$31*{c}{ADR}', '#,##0', True, 'Target × adoption'),
    ('Share-of-wallet uplift (pp)', lambda c: f'={c}{SHR}-{base}', '0.0%', False, 'Adopter share − today'),
    ('Incremental GMV per month (₹ Cr)', lambda c: f'={c}37*{spend}*{c}38/10000000', '#,##0', False, ''),
    ('Incremental GMV per year (₹ Cr)', lambda c: f'={c}39*12', '#,##0', True, 'Main volume number'),
    ('Incremental UPI transactions per month', lambda c: f'={c}37*{bills}*{c}38', '#,##0', False, ''),
    ('Paytm share of segment supplier UPI', lambda c: f'={base}+{c}39/$B$32', '0.0%', False, 'MAIN METRIC (slide 13)'),
    ('Added to Paytm share of all UPI value (pp)', lambda c: f'={c}39/{nat}*100', '0.000', False, ''),
    ('Revenue: payer-app MDR', lambda c: f'={c}40*{elig}*{tpap}', '0.0', True, ''),
    ('Revenue: acquirer MDR', lambda c: f'={c}40*{elig}*{acqs}*{acq}', '0.0', True, ''),
    ('Credit users', lambda c: f'={c}37*{c}{CUR}', '#,##0', False, ''),
    ('Credit disbursed (₹ Cr / yr)', lambda c: f'={c}46*{camt}*12/10000000', '#,##0', True, ''),
    ('Revenue: credit distribution', lambda c: f'={c}47*{take}', '0.0', True, ''),
    ('Total revenue', lambda c: f'={c}44+{c}45+{c}48', '0.0', True, ''),
    ('New active shops in year', lambda c: '=B37' if c == 'B' else f'={c}37-{chr(ord(c)-1)}37', '#,##0', True, ''),
    ('Cost: build + run', lambda c: f'={build}+{run}' if c == 'B' else f'={run}', '0.0', True, ''),
    ('Cost: acquisition', lambda c: f'={c}50*{cac}/10000000', '0.0', True, ''),
    ('Cost: AI bill reading', lambda c: f'={c}37*{bills}*12*{ocr}*{cap}/10000000', '0.0', True, ''),
    ('Cost: support', lambda c: f'={c}37*{sup}/10000000', '0.0', True, ''),
    ('Total cost', lambda c: f'=SUM({c}51:{c}54)', '0.0', True, ''),
    ('Contribution', lambda c: f'={c}49-{c}55', '0.0', True, ''),
    ('Cost per extra active shop (₹)', lambda c: f'={c}55*10000000/{c}37', '#,##0', False, 'Brief metric'),
    ('Cost per ₹1,000 of extra GMV (₹)', lambda c: f'={c}55/{c}40*1000', '0.00', False, 'Brief metric'),
]
for i, (k, fn, fm, tot, note) in enumerate(lines, 37):
    bold = k.startswith(('Total', 'Contribution', 'Paytm share', 'Incremental GMV per year'))
    put(bm, f'A{i}', k, fBold if bold else fB)
    for c in 'BCD':
        put(bm, f'{c}{i}', fn(c), fBold if bold else fB, fmt=fm, al=Alignment(horizontal='right'))
    if tot:
        put(bm, f'E{i}', f'=SUM(B{i}:D{i})', fB, fmt=fm, al=Alignment(horizontal='right'))
    put(bm, f'F{i}', note, fNote)
r = 60
header(bm, r, ['E. Per-shop economics (Year 3 share)', 'Value', '', '', '', 'Note'])
pe = [('New Paytm GMV per active shop per year (₹)', f'={spend}*12*(D{SHR}-{base})', '#,##0', ''),
      ('Revenue per shop: payer-app MDR (₹)', f'=B61*{elig}*{tpap}', '#,##0', ''),
      ('Revenue per shop: acquirer MDR (₹)', f'=B61*{elig}*{acqs}*{acq}', '#,##0', ''),
      ('Revenue per shop: credit (₹)', f'=D{CUR}*{camt}*12*{take}', '#,##0', ''),
      ('CAC payback, payments only (months)', f'={cac}/((B62+B63)/12)', '0.0', ''),
      ('CAC payback incl. credit (months)', f'={cac}/((B62+B63+B64)/12)', '0.0', ''),
      ('3-yr LTV / CAC, payments only', f'=(B62+B63)*(1+0.7+0.49)/{cac}', '0.0"×"', '70% annual retention, undiscounted'),
      ('3-yr LTV / CAC incl. credit', f'=(B62+B63+B64)*(1+0.7+0.49)/{cac}', '0.0"×"', ''),
      ('Comparison: 0.1% cashback on Year 3 GMV (₹ Cr / yr)', '=D40*0.001', '0.0', 'Recurring every year, vs total Year 3 cost in D55')]
for i, (k, f, fm, note) in enumerate(pe, 61):
    put(bm, f'A{i}', k); put(bm, f'B{i}', f, fBold, fmt=fm, al=Alignment(horizontal='right')); put(bm, f'F{i}', note, fNote)
bm.freeze_panes = 'B5'

# ======================= SENSITIVITY =======================
se = sheet('Sensitivity', 'Sensitivity · Year 3, one assumption at a time (slide 15)',
           'Low / high values are blue inputs. Other inputs are linked (green) to Business Model.', [46, 13, 13, 18, 18, 18, 18])
BM = "'Business Model'!"
put(se, 'A4', 'Base Year 3 incremental GMV (₹ Cr / yr)'); put(se, 'B4', f"={BM}D40", fLink, fmt='#,##0')
put(se, 'A5', 'Base Year 3 revenue (₹ Cr / yr)'); put(se, 'B5', f"={BM}D49", fLink, fmt='0.0')
header(se, 7, ['Driver', 'Low', 'High', 'GMV at low', 'GMV at high', 'Revenue at low', 'Revenue at high'])
T, AD, SH, SP, BS, EL, TP, AS, AC, CU, CA, TK = (f"{BM}$B$31", f"{BM}$D${ADR}", f"{BM}$D${SHR}", f"{BM}$B$8", f"{BM}$B$9", f"{BM}$B$11",
                                                  f"{BM}$B$12", f"{BM}$B$13", f"{BM}$B$14", f"{BM}$D${CUR}", f"{BM}$B$15", f"{BM}$B$16")
def g(ad=AD, sp=SP, sh=SH): return f"{T}*{ad}*{sp}*({sh}-{BS})*12/10000000"
def rv(ad=AD, sp=SP, sh=SH, el=EL, cu=CU):
    return f"{g(ad, sp, sh)}*{el}*({TP}+{AS}*{AC})+{T}*{ad}*{cu}*{CA}*12/10000000*{TK}"
sens = [('Year 3 adoption of target shops', 0.20, 0.50, '0%', lambda x: dict(ad=x)),
        ('Supplier UPI per shop per month (₹)', 40000, 80000, '#,##0', lambda x: dict(sp=x)),
        ('Paytm share of an adopter\'s supplier UPI', 0.55, 0.85, '0%', lambda x: dict(sh=x)),
        ('MDR-eligible share of Chukta GMV', 0.20, 0.70, '0%', lambda x: dict(el=x)),
        ('Credit users among adopters (Year 3)', 0.10, 0.30, '0%', lambda x: dict(cu=x))]
for i, (k, lo, hi, fm, kw) in enumerate(sens, 8):
    put(se, f'A{i}', k); put(se, f'B{i}', lo, fIn, fmt=fm); put(se, f'C{i}', hi, fIn, fmt=fm)
    for col, ref, fun, ff in [('D', f'B{i}', g, '#,##0'), ('E', f'C{i}', g, '#,##0'), ('F', f'B{i}', rv, '0.0'), ('G', f'C{i}', rv, '0.0')]:
        args = kw(ref)
        if fun is g: args = {k2: v2 for k2, v2 in args.items() if k2 in ('ad', 'sp', 'sh')}
        put(se, f'{col}{i}', '=' + fun(**args), fB, fmt=ff)
put(se, 'A14', 'Reading: adoption swings GMV most; credit uptake swings revenue most. GMV stays above ₹37,000 Cr in every tested case.', fNote)

# ======================= PILOT DESIGN =======================
pd_ = sheet('Pilot Design', 'Pilot Design · 90-day cluster-randomised Kanpur pilot (slide 12)', None, [46, 16, 60])
header(pd_, 4, ['Element', 'Value', 'Detail'])
rows(pd_, 5, [
    ('Clusters (distributors)', '20', 'FMCG + pharma, matched in pairs by size and category; one of each pair randomised to treatment'),
    ('Treatment / control', '10 / 10', '~2,500 retailers per arm; ~1,500 Paytm device merchants per arm form the analysis cohort'),
    ('Baseline', '4 weeks', 'Distributor ledgers (invoice, amount, payment mode, app) + Paytm data; Account Aggregator consent on a 200-shop subsample'),
    ('Primary metric', 'DiD', 'Paytm share of supplier ₹ (treated − control, post − pre), standard errors clustered by distributor'),
    ('Credit', 'Week 7+', 'Lender-approved sub-cohort only, analysed separately'),
    ('Incentive holdout', 'Day 60', 'Distributor early-pay discount switched off for a random half of treated distributors'),
    ('Budget', '₹10 L', 'Field team 6 × 3 months ₹7.2 L · activation fees ₹0.5 L · onboarding ₹0.5 L · infra ₹0.3 L · comms ₹0.5 L · contingency ₹1 L')])
header(pd_, 13, ['Power calculation', 'Value', 'Formula / note'])
pw = [('Analysable shops per arm', 1500, None, 'Input'), ('Clusters per arm', 10, None, 'Input'), ('Intra-cluster correlation (ICC)', 0.05, None, 'Assumed'),
      ('SD of shop-level Paytm share', 0.35, None, 'Assumed (share is 0–1, bimodal)'), ('Shops per cluster', None, '=B14/B15', ''),
      ('Design effect', None, '=1+(B18-1)*B16', '1 + (m − 1) × ICC'), ('Effective shops per arm', None, '=B14/B19', ''),
      ('Minimum detectable effect (80% power, α = 0.05)', None, '=(1.96+0.84)*B17*SQRT(2/B20)', 'Target effect is +25 pp')]
for i, (k, v, f, note) in enumerate(pw, 14):
    put(pd_, f'A{i}', k)
    if v is not None: put(pd_, f'B{i}', v, fIn, fmt='0.00' if v < 1 else '#,##0')
    else: put(pd_, f'B{i}', f, fBold, fmt='0.0%' if 'detectable' in k else '0.0')
    put(pd_, f'C{i}', note, fNote)
header(pd_, 23, ['Day-90 target (treated Paytm device shops)', 'Target', 'Decision use'])
rows(pd_, 24, [('Paytm share of supplier ₹: DiD vs control', '≥ +25 pp', 'Primary scale criterion'),
               ('Paytm share of supplier ₹ in treated shops', '≥ 35%', ''), ('Activation: ≥ 1 bill paid in 14 days', '≥ 40%', 'STOP if < 15%'),
               ('Repeat: ≥ 3 bills in month 3 (of activated)', '≥ 60%', 'Scale criterion'), ('Invoices visible in Chukta', '≥ 70%', ''),
               ('Retention after discount switched off at day 60', '≥ 85%', 'Incentive-independence test'), ('Distributor days-to-collect', '−3 days', ''),
               ('Guardrails', 'Failures < 1% · complaints < 2% · churn ≤ control', 'Must hold to scale')])
header(pd_, 33, ['Decision rule', 'Condition', 'Next step'])
rows(pd_, 34, [('SCALE', 'DiD ≥ +25 pp and repeat ≥ 60% and guardrails met', '10 high-Soundbox cities, 500 distributors, credit live'),
               ('ITERATE', 'DiD +10 to +25 pp', 'Fix weakest funnel step; extend 60 days'),
               ('STOP', 'DiD < +10 pp or activation < 15%', 'Keep only the Soundbox reminder')])

# ======================= JOURNEY / BUILD / RISKS =======================
jb = sheet('Journey · Build · Risks', 'Before / after journey, build plan and risk register (slides 8, 10, 15)', None, [30, 40, 22, 14, 50])
header(jb, 4, ['BEFORE step', 'What happens', 'Pain', '', ''])
before = [('1 Bill arrives', 'Paper challan or WhatsApp PDF', 'No due date tracked'), ('2 Filed', 'Drawer or chat', 'Lost / forgotten'),
          ('3 Collection', 'Salesman visits or calls', 'Owner on the spot'), ('4 App', 'PhonePe / GPay on personal phone', 'Paytm never in hand'),
          ('5 QR', 'Scans whichever QR the salesman shows', 'Salesman picks app'), ('6 Amount', 'Typed from challan or memory', 'Typos, part-payments'),
          ('7 Proof', 'Screenshot on WhatsApp', 'No invoice link'), ('8 Matching', 'Accountant matches by hand', 'Disputes'),
          ('9 Cash short', 'Udhaar or delay', 'Strains credit terms')]
rows(jb, 5, before)
header(jb, 15, ['AFTER step', 'What happens', 'Paytm asset', '', ''])
rows(jb, 16, [('1 Bill lands', 'Distributor push / e-invoice QR / photo', 'AI read, IRP QR verify'),
              ('2 Reminder', 'Soundbox 9 pm: tomorrow\'s dues', 'Soundbox summary voice'),
              ('3 Pay', 'Payee verified; invoice-tagged UPI; or 7 din baad', 'Paytm UPI, credit line on UPI'),
              ('4 Close', 'Distributor auto-reconciled; purchase register', 'Chukta Collect')])
header(jb, 21, ['Component', 'New / reuse', 'Effort (S/M/L)', 'In MVP?', 'Note'])
rows(jb, 22, [('Bill inbox + ledger + due-date engine', 'New', 'M', 'Yes', 'Both apps, synced by merchant ID'),
              ('e-invoice QR read + IRP signature check', 'New', 'S', 'Yes', 'Public schema and key'),
              ('Photo / PDF bill read', 'New (Paytm AI)', 'M', 'Yes', 'Owner confirms'),
              ('Payee verification GSTIN ↔ VPA / a/c', 'New', 'S', 'Yes', 'Flags mismatches'),
              ('Soundbox "due tomorrow" line', 'Reuse', 'S', 'Yes', 'Script extension'),
              ('Invoice-tagged UPI pay', 'Reuse', 'S', 'Yes', 'Rail picked by supplier'),
              ('Chukta Collect web view', 'New', 'M', 'Yes', 'Read-only first'),
              ('Distributor push (API, Bharat Connect, Tally/Marg/Busy)', 'New + partners', 'L', 'CSV in pilot', ''),
              ('Business credit line on UPI', 'Partner', 'L', 'Week 7 sub-cohort', 'P2M payees only'),
              ('GST purchase register export', 'New', 'S', 'Phase 2', '')])
header(jb, 33, ['Risk', 'Mitigation', 'Early signal', 'L / I', ''])
rows(jb, 34, [('PhonePe / GPay copy the bill inbox', 'Distributor lock-in first; ERP buttons; two-sided credit data', 'Competitor launch; distributor churn', 'H / M'),
              ('Competitor cashback war', 'Don\'t match; workflow + records + credit', 'Treated share falls > 5 pp', 'M / M'),
              ('Distributors avoid MDR', 'Rail-neutral: bank-a/c rail is MDR-free', 'Rail mix in ledgers', 'H / L'),
              ('Fake bill / QR-swap fraud', 'GSTIN name match; e-invoice signature; 24 h cooling-off; push-only', 'Name-mismatch rate', 'M / H'),
              ('Credit losses', 'Lender underwrites (RBI DLD 2025); small limits', 'First-payment default', 'M / M'),
              ('Habit doesn\'t form', 'Nightly Soundbox ritual; STOP rule', 'Repeat < 40% at D60', 'M / H'),
              ('Rule change (MDR rollback)', 'MDR is a minority of revenue, 0% of GMV case', 'NPCI circulars', 'M / L'),
              ('Data privacy (DPDP)', 'Consent per use; delete photos after reading', 'Opt-out rate', 'L / M')])

# ======================= R1 VOC (values copy) =======================
import os, glob
_here = os.path.dirname(os.path.abspath(__file__))
_cands = [os.path.join(_here, 'inputs', 'KartikRaj_VOC.xlsx')] + glob.glob('/mnt/user-data/uploads/*KartikRaj_VOC.xlsx')
_vp = next((c for c in _cands if os.path.exists(c)), None)
if not _vp: raise SystemExit('Put Round 1 KartikRaj_VOC.xlsx in ./inputs/ first')
src = load_workbook(_vp, data_only=True)['VOC Collection']
rv_ = wb.create_sheet('R1 VOC'); rv_.sheet_view.showGridLines = False
rv_['A1'] = 'Round 1 VOC Collection (50 respondents) · values copied from KartikRaj_VOC.xlsx'; rv_['A1'].font = fT
hdr_row = None
for row in src.iter_rows(min_row=1, max_row=src.max_row):
    if str(row[0].value).strip() == 'VOC ID': hdr_row = row[0].row; break
out = 3
for rr in range(hdr_row, src.max_row + 1):
    vals = [src.cell(row=rr, column=c).value for c in range(1, src.max_column + 1)]
    if rr > hdr_row and not str(vals[0] or '').startswith('VOC'): continue
    for c, v in enumerate(vals, 1):
        cell = rv_.cell(row=out, column=c, value=v); cell.border = bd; cell.alignment = wrap
        cell.font = fH if rr == hdr_row else fB
        if rr == hdr_row: cell.fill = fillH
    out += 1
for c in range(1, src.max_column + 1): rv_.column_dimensions[L(c)].width = 16
rv_.freeze_panes = 'B4'

# ======================= SOURCES =======================
so = sheet('Sources', 'Sources · every external number used in the deck', None, [64, 95])
header(so, 3, ['Fact used', 'Source'])
rows(so, 4, [
    ('Paytm Q1 FY27: 1.57 Cr device merchants, merchant GMV ₹7.1 L Cr (+31%), FS revenue ₹814 Cr (+45%), MTU 8 Cr', 'https://upstox.com/news/market-news/earnings/paytm-q1-results-revenue-jumps-28-yo-y-to-2-448-crore-ebitda-zooms-182-yo-y-check-key-metrics/article-197256/'),
    ('Paytm Q1 FY27 detail (device merchants 1.57 Cr; MTU)', 'https://hdfcsky.com/news/paytm-q1-fy27-profit-soars-79percent-ebitda-hits-record-high'),
    ('UPI MDR 0.4% on P2M > ₹2,000 from 15 Oct 2026; ₹300 cap; ≤ ₹2,000 free', 'https://www.angelone.in/news/economy/npci-introduces-0-4-mdr-on-upi-payments-above-2-000-from-october-15-2026'),
    ('P2PM small merchants (≤ ₹1 L/month) exempt; MDR borne by merchants', 'https://www.icicidirect.com/research/equity/blog/npci-will-introduce-a-mdr-on-select-p2m-upi-transactions-above-rs-2000-from-october-15'),
    ('MDR split: issuer 40%, acquirer 30%, UPI apps 20%, app bank partners 10%', 'https://india.entrepreneur.com/business-news/npci-announces-0-4-mdr-on-upi-transactions-above-inr-2000'),
    ('NPCI MDR FAQs, 15 Sep 2026', 'https://www.scconline.com/blog/post/2026/09/16/npci-released-upi-mdr-faqs-explained/'),
    ('UPI P2P collect discontinued from 1 Oct 2025 (circular 29 Jul 2025)', 'https://www.outlookmoney.com/banking/npci-to-end-upi-p2p-collect-requests-from-october-1-to-reduce-fraud'),
    ('NPCI 30% cap due 31 Dec 2026; PhonePe + GPay 78.6% of volume (Jul 2026)', 'https://startupfeed.in/upi-30-percent-cap-deadline-phonepe-google-pay-2026/'),
    ('UPI app shares 2026 (PhonePe ~46%, GPay ~32.7%, Paytm ~7.9%)', 'https://currentaffairs.adda247.com/top-10-upi-apps-in-india-by-market-share-in-2026/'),
    ('Paytm UPI share Aug 2026 8.1% (R1 source)', 'https://inc42.com/buzz/upi-in-august-navis-market-share-climbs-to-4-4-phonepe-google-pay-slip/'),
    ('Paytm AI Soundbox launch (9 Oct 2025; 11 languages; voice assistant)', 'https://paytm.com/blog/artificial-intelligence/paytm-launches-indias-first-ai-soundbox-for-payments/'),
    ('Paytm for Business app (10 M+ businesses; receiving-focused)', 'https://play.google.com/store/apps/details?id=com.paytm.business&hl=en_IN'),
    ('PhonePe Business app (5 Cr merchants; speakers; EDI loans; no supplier-bill feature listed)', 'https://play.google.com/store/apps/details?id=com.phonepe.app.business&hl=en_IN'),
    ('Paytm Postpaid credit line on UPI with Suryoday SFB (merchant payments only)', 'https://www.medianama.com/2025/09/223-paytm-upi-credit-line-suryoday-bank/'),
    ('Credit line on UPI: banks offering; SFBs permitted', 'https://www.business-standard.com/amp/finance/news/rbi-allows-small-finance-banks-to-offer-pre-sanctioned-credit-line-on-upi-124120600474_1.html'),
    ('Bharat Connect for Business: 2 mn+ invoice txns Q3 FY26, ~50% MoM growth', 'https://deshgujarat.com/2026/02/26/nbbls-bharat-connect-for-business-emerges-as-digital-backbone-for-gujarats-43-lakh-msmes/'),
    ('TallyPrime 7.0 integrates Bharat Connect for Business', 'https://tallysolutions.com/tally/b2b-bharat-connect-tallyprime/'),
    ('E-invoice: ₹5 Cr threshold; signed QR fields (GSTINs, invoice no., date, value, IRN)', 'https://busy.in/gst/invoice-registration-portal-irp-under-gst.md'),
    ('E-invoice threshold unchanged in 2026', 'https://tallysolutions.com/accounting/e-invoicing-rules-in-india/'),
    ('FMCG distributor margins 4–8% GT; retailer credit 7–21 days', 'https://spirestock.com/blog/fmcg-distributor-margin-profit-guide-india'),
    ('Kirana count ~13 M; distributor credit across hundreds of small accounts', 'https://spirestock.com/blog/kirana-store-meaning-fmcg-distribution'),
    ('Bizom: 600+ brands, 8 M retailers', 'https://happierleads.com/companies/bizom-com'),
    ('Marg ERP 6 L+ users', 'https://www.softwaresuggest.com/us/distribution-management-software'),
    ('Round 1 primary research: 50 VOCs, Kanpur, Sep 2026', 'KartikRaj_VOC.xlsx (sheet "R1 VOC" in this workbook)')])

os.makedirs(os.path.join(_here, 'out'), exist_ok=True)
wb.save(os.path.join(_here, 'out', 'KartikRaj_R2_Annexure.xlsx'))
print('saved')
