"""Copy field answers from the user's working file into out/KartikRaj_R2_Annexure.xlsx (rows 6+)."""
import sys, datetime
import pandas as pd
from openpyxl import load_workbook
src, dst = sys.argv[1], sys.argv[2]
DATE = datetime.datetime(2026, 10, 8)
m = pd.read_excel(src, 'R2 Validation Log', keep_default_na=False); d = pd.read_excel(src, 'Distributor Log', keep_default_na=False)
route = {'e-Invoice QR': 'e-Invoice QR', 'Photo/PDF/WhatsApp': 'Paper / photo / WhatsApp'}
wb = load_workbook(dst)
ws = wb['R2 Validation Log']
for i, r in m.iterrows():
    row = [r['Respondent ID'], DATE, r['Area'], r['Category'], r['Role'], r['Receives on'], r['Round 1 shop?'],
           int(r['Last supplier payment']), r['App used'], r['Paid from'], int(r['Minutes taken']), int(r['A1 Useful']),
           r['A2 Forced choice'], r['A3 Verbatim'], r['A3 Objection code'], route[r['T1 Route']],
           r['L1 Data'], r['L2 Time'], r['L3 Money'], r['L4 Promise'], r['Personal phone'],
           r['Notes'] or None]
    for j, v in enumerate(row, 1): ws.cell(row=6 + i, column=j, value=v)
ws2 = wb['Distributor Log']
mdr = {'Yes (concerned)': 'Yes (concerned)', 'No/Not aware': 'No / not aware'}
erp = {'None or paper': 'None / paper'}
for i, r in d.iterrows():
    row = [r['Respondent ID'], DATE, r['Area'], r['Category'], int(r['Retailers served']), r['Monthly sales'], int(r['Credit days']),
           erp.get(r['ERP software'], r['ERP software']), r['e-Invoice QR?'], int(r['Cash %']), int(r['UPI QR %']), int(r['Bank transfer %']),
           mdr[r['MDR aware?']], r['Ledger shared?'], r['LOI signed?'], r['Key quote'],
           r['Notes'] or None]
    for j, v in enumerate(row, 1): ws2.cell(row=6 + i, column=j, value=v)
wb.save(dst); print(len(m), 'merchants,', len(d), 'distributors')
