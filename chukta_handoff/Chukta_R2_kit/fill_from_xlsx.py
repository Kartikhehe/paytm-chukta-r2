"""After fieldwork: python3 fill_from_xlsx.py path/to/KartikRaj_R2_Annexure.xlsx
Reads the 'R2 Validation Dashboard' (column D = deck key, column B = value) into validation.json,
then rebuilds the PDF. Open + save the workbook in Excel (or run LibreOffice recalc) first so
formula values are cached. Objection -> design-change rows are edited by hand in validation.json.
"""
import json, sys, subprocess, os
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
xl = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, 'KartikRaj_R2_Annexure.xlsx')
ws = load_workbook(xl, data_only=True)['R2 Validation Dashboard']
vpath = os.path.join(HERE, 'validation.json')
v = json.load(open(vpath))
found = 0
pairs = []
for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
    pairs.append((row[3].value, row[1].value))
    if len(row) > 7: pairs.append((row[7].value, row[6].value))
for key, val in pairs:
    if key and key in v and isinstance(val, (int, float)):
        v[key] = int(round(val)) if float(val).is_integer() or str(key).endswith('_s') else round(val, 1)
        found += 1
if not v.get('merchants_tested'):
    sys.exit('No respondents found in the log. Fill the R2 Validation Log, save in Excel, then rerun.')
json.dump(v, open(vpath, 'w'), indent=2, ensure_ascii=False)
print(f'updated {found} fields in validation.json')
print('Still blank:', [k for k, x in v.items() if x is None and not k.startswith('_')])
print('Remember to fill heard_changed, test_dates, proto_link and video_link by hand.')
subprocess.run([sys.executable, os.path.join(HERE, 'build.py'), '--png'], check=True)
